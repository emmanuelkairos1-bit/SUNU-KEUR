# -*- coding: utf-8 -*-
"""Contrôle qualité PPTX : estime si le texte tient dans ses cadres (heuristique).
Signale les débordements probables (lignes nécessaires vs hauteur du cadre)."""
import math
from pptx import Presentation
from pptx.util import Emu

PPTX = r"C:\dev\SUNU-KEUR\docs\SUNU-KEUR-Guide-Fonctionnement.pptx"
FACTEUR = 0.50   # largeur moyenne d'un caractère / corps (mixte gras+normal)
TOLERANCE = 1.22  # marge autorisée sur la hauteur


def pt(v):
    return Emu(v).pt if v is not None else 0.0


def controler():
    prs = Presentation(PPTX)
    problemes = []
    for num, slide in enumerate(prs.slides, 1):
        for shp in slide.shapes:
            if not shp.has_text_frame:
                continue
            tf = shp.text_frame
            texte = tf.text
            if not texte.strip():
                continue
            larg = pt(shp.width) - pt(tf.margin_left) - pt(tf.margin_right)
            haut = pt(shp.height) - pt(tf.margin_top) - pt(tf.margin_bottom)
            if larg <= 8:
                continue
            besoin_h = 0.0
            depasse_l = False
            for para in tf.paragraphs:
                corps = max([(r.font.size.pt if r.font.size else 11) for r in para.runs] or [11])
                gras = any(r.font.bold for r in para.runs)
                txt = "".join(r.text for r in para.runs)
                if not txt:
                    besoin_h += corps * 1.2
                    continue
                largeur_est = len(txt) * corps * (FACTEUR + (0.03 if gras else 0))
                if tf.word_wrap is False:
                    if largeur_est > larg:
                        depasse_l = True
                    nlignes = 1
                else:
                    nlignes = max(1, math.ceil(largeur_est / larg))
                espacement = para.line_spacing or 1.0
                besoin_h += nlignes * corps * 1.22 * (espacement if espacement > 0.2 else 1.0)
                besoin_h += (para.space_before.pt if para.space_before else 0)
                besoin_h += (para.space_after.pt if para.space_after else 0)
            ratio = besoin_h / haut if haut > 0 else 99
            if ratio > TOLERANCE:
                problemes.append(
                    f"S{num:02d} TEXTE TROP LONG x{ratio:.2f} (h cadre {haut:.0f} vs besoin {besoin_h:.0f}) "
                    f"« {texte[:46]} »"
                )
            if depasse_l:
                problemes.append(f"S{num:02d} TEXTE HORS LARGEUR « {texte[:46]} »")
    if problemes:
        print(f"=== {len(problemes)} ALERTES FITTING ===")
        for p in problemes:
            print(" ", p)
    else:
        print("=== FITTING OK — tous les textes tiennent dans leurs cadres ===")
    return len(problemes)


if __name__ == "__main__":
    raise SystemExit(1 if controler() else 0)
